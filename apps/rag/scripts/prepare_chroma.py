import json
import logging
import os
from pathlib import Path
from tqdm import tqdm
from PIL import Image
import pytesseract

from langchain_community.document_loaders import TextLoader, NotebookLoader, PyPDFLoader
from langchain.schema import Document

from apps.rag.utils import load_embedding_function, get_chroma_collection_native
from apps.rag.splitter import get_splitter
from apps.rag.module_loader import module_loader

logger = logging.getLogger(__name__)

# La carte « répertoire de module → fichier d'index ».
#
# Compétence visée : C18 (épreuve E4)
# Elle venait de `apps.rag.module_index_map`, un module de compatibilité qui
# ne faisait que relayer `module_loader` et qui a été supprimé le 30/08/2026
# avec le remplacement de la liste blanche maison par RestrictedPython. Cet
# import est resté, si bien que ce script ne s'importait plus du tout —
# `ModuleNotFoundError` — sans que rien ne le signale, puisque rien ne
# l'importe. Il est branché directement sur sa source au lieu de passer par
# un relais.
MODULE_INDEX_MAP = module_loader.module_index_map

# === CONFIGURATION ===
DATA_FOLDER = Path("data/contents")
COURSES_FOLDER = DATA_FOLDER / "courses"
INDEX_FOLDER = DATA_FOLDER / "index"
RESOURCES_FOLDER = DATA_FOLDER / "resources"
CHUNK_THRESHOLD = 1000

SUPPORTED_IMAGE_EXTS = [".png", ".jpg", ".jpeg", ".avif"]
SUPPORTED_TEXT_EXTS = [".md", ".ipynb", ".pdf", ".pptx"]

_loaded_indexes = {}

# === SECTION DETECTION LOGIC ===
def get_section(filepath: Path):
    filename = filepath.name
    parent_dir = filepath.parent.name

    # --- Course case ---
    if parent_dir in MODULE_INDEX_MAP:
        index_file = MODULE_INDEX_MAP[parent_dir]
        if index_file not in _loaded_indexes:
            try:
                with open(INDEX_FOLDER / index_file, "r", encoding="utf-8") as f:
                    _loaded_indexes[index_file] = json.load(f)
            except Exception as e:
                print(f"❌ Error reading {index_file}: {e}")
                return "unknown"

        index = _loaded_indexes[index_file]
        for section, files in index.items():
            if filename in files:
                return section
        return "unknown"

    # --- Resources case ---
    elif parent_dir == "resources":
        index_file = "ressources_index.json"
        if index_file not in _loaded_indexes:
            try:
                with open(INDEX_FOLDER / index_file, "r", encoding="utf-8") as f:
                    _loaded_indexes[index_file] = json.load(f)
            except Exception as e:
                print(f"❌ Error reading {index_file}: {e}")
                _loaded_indexes[index_file] = {}

        index = _loaded_indexes[index_file]
        for section, files in index.items():
            if filename in files:
                return section

        # Fallback by name
        if "__" in filename:
            return filename.split("__")[0]

        return "ressources"

    return "unknown"

# === OCR for images ===
def ocr_image_to_document(filepath: Path):
    try:
        text = pytesseract.image_to_string(Image.open(filepath))
        return Document(
            page_content=text,
            metadata={
                "source": filepath.name,
                "type": "image_cheatsheet",
                "section": get_section(filepath)
            }
        )
    except Exception as e:
        print(f"❌ OCR failed for {filepath.name}: {e}")
        return None

# === Diaporamas ===
def pptx_to_document(filepath: Path):
    """
    Rend le texte d'un diaporama, une diapositive par bloc.

    Compétence visée : C10 (épreuve E3), C4 (E1)

    Choix : `python-pptx` plutôt que `UnstructuredPowerPointLoader`.
    Motivation : ce dernier tire `unstructured`, qui entraîne à sa suite une
    chaîne de dépendances d'analyse documentaire hors de proportion avec le
    besoin. `python-pptx` lit le format directement, et la seule chose qu'on
    lui demande est le texte.

    Choix : les diapositives sont séparées par un titre de niveau 2 numéroté.
    Motivation : un diaporama concaténé d'un bloc perd sa structure, et le
    découpeur du RAG couperait au milieu d'une notion. Le numéro de diapositive
    donne au fragment une adresse que l'apprenant peut retrouver dans le
    fichier d'origine.

    Choix : le texte des tableaux est repris, cellule par cellule. Motivation :
    dans un support sur les relations entre tables, le tableau EST le contenu ;
    l'ignorer laisserait des diapositives vides.
    """
    from pptx import Presentation

    presentation = Presentation(str(filepath))
    morceaux = []

    for rang, diapositive in enumerate(presentation.slides, start=1):
        lignes = []
        for forme in diapositive.shapes:
            if forme.has_text_frame and forme.text_frame.text.strip():
                lignes.append(forme.text_frame.text.strip())
            if getattr(forme, "has_table", False):
                for ligne in forme.table.rows:
                    cellules = [c.text.strip() for c in ligne.cells if c.text.strip()]
                    if cellules:
                        lignes.append(" | ".join(cellules))
        if lignes:
            morceaux.append(f"## Diapositive {rang}\n\n" + "\n\n".join(lignes))

    if not morceaux:
        logger.warning("Diaporama sans texte exploitable : %s", filepath.name)
        return None

    return Document(
        page_content="\n\n".join(morceaux),
        metadata={
            "source": filepath.name,
            "type": "pptx",
            "section": get_section(filepath),
        },
    )


# === Unit loader ===
def load_document(filepath: Path):
    suffix = filepath.suffix.lower()

    try:
        if suffix == ".md":
            return TextLoader(str(filepath), encoding="utf-8").load()
        elif suffix == ".ipynb":
            return NotebookLoader(str(filepath)).load()
        elif suffix == ".pdf":
            return PyPDFLoader(str(filepath)).load()
        elif suffix == ".pptx":
            doc = pptx_to_document(filepath)
            return [doc] if doc else []
        elif suffix in SUPPORTED_IMAGE_EXTS:
            doc = ocr_image_to_document(filepath)
            return [doc] if doc else []
        else:
            raise ValueError(f"Unsupported file type: {suffix}")
    except Exception as e:
        print(f"❌ Error loading {filepath.name}: {e}")
        return []

# === Complete folder indexing ===
def process_directory(path: Path, collection, splitter, file_type: str):
    for file in tqdm(path.rglob("*"), desc=f"Indexing {file_type}"):
        if not file.is_file():
            continue
        if file.suffix.lower() not in SUPPORTED_TEXT_EXTS + SUPPORTED_IMAGE_EXTS:
            continue

        docs = load_document(file)
        if not docs:
            continue

        section = get_section(file)

        if file_type == "resources":
            existing_sections = {
                s.lower().split("_", 1)[1]  # "01_python" → "python"
                for s in MODULE_INDEX_MAP.keys()
            }
            if section not in existing_sections:
                print(f"⏩ Skip {file.name} (section '{section}' not yet covered)")
                continue

        for doc in docs:
            doc.metadata.update({
                "source": file.name,
                "type": file.suffix[1:],  # "md", "pdf"...
                "section": section
            })

        chunks = docs if len(docs[0].page_content) < CHUNK_THRESHOLD else splitter.split_documents(docs)

        try:
            collection.add(
                documents=[chunk.page_content for chunk in chunks],
                metadatas=[chunk.metadata for chunk in chunks],
                ids=[f"{file.stem}-{i}" for i in range(len(chunks))]
            )
        except Exception as e:
            print(f"❌ Failed to index {file.name}: {e}")
            
        print(f"✅ {file.name} → {len(chunks)} chunk(s)")

# === ENTRY POINT ===
def main():
    collection = get_chroma_collection_native()
    splitter = get_splitter()

    for module_dir in MODULE_INDEX_MAP.keys():
        full_path = COURSES_FOLDER / module_dir
        if full_path.exists():
            process_directory(full_path, collection, splitter, "courses")
        else:
            print(f"⚠️ Folder {full_path} not found, ignored.")

    if RESOURCES_FOLDER.exists():
        process_directory(RESOURCES_FOLDER, collection, splitter, "resources")

    print("✅ Chroma vectorstore fully populated.")

if __name__ == "__main__":
    main()
